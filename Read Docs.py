from docx import Document
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
import glob, os
from tkinter import *
import tkinter as tk
from tkinter import messagebox


def first_button():
    working_directory = e1.get()
    folder_directory = e2.get()

    try:

        os.chdir(folder_directory)

        for file in glob.glob("*.docx"):

            os.chdir(folder_directory)

            wordDoc = Document(file)

            InputDictionary = []

            for table in wordDoc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        InputDictionary.append(cell.text)

            InputResultsTable = {}
            InputResultsTable["Evidence Number"] = InputDictionary[1]
            InputResultsTable["Case Information - Date"] = InputDictionary[10]
            InputResultsTable["Case Information - Custodian"] = InputDictionary[11]
            InputResultsTable["Case Information - Matter Name"] = InputDictionary[12]
            InputResultsTable["Case Information - Matter Number"] = InputDictionary[13]
            InputResultsTable["Case Information - Address Location"] = InputDictionary[15]
            InputResultsTable["Type of ESI - Received By"] = InputDictionary[32]
            InputResultsTable["Type of ESI - Desktop"] = InputDictionary[45]
            InputResultsTable["Type of ESI - Laptop"] = InputDictionary[47]
            InputResultsTable["Type of ESI - Server"] = InputDictionary[49]
            InputResultsTable["Type of ESI - Tablet"] = InputDictionary[51]
            InputResultsTable["Type of ESI - Mobile Device"] = InputDictionary[53]
            InputResultsTable["Type of ESI - Hard Disk Drive"] = InputDictionary[55]
            InputResultsTable["Type of ESI - Optical Disc"] = InputDictionary[58]
            InputResultsTable["Type of ESI - Removable Media"] = InputDictionary[60]
            InputResultsTable["Type of ESI - Removable Media Other"] = InputDictionary[64]
            InputResultsTable["Type of ESI - Additional Details"] = InputDictionary[71]
            InputResultsTable["Details of ESI - Device Make"] = InputDictionary[88]
            InputResultsTable["Details of ESI - Media Make"] = InputDictionary[90]
            InputResultsTable["Details of ESI - Device Model"] = InputDictionary[92]
            InputResultsTable["Details of ESI - Media Model"] = InputDictionary[94]
            InputResultsTable["Details of ESI - Device Serial"] = InputDictionary[96]
            InputResultsTable["Details of ESI - Media Serial"] = InputDictionary[98]
            InputResultsTable["Details of ESI - Device Asset Tag"] = InputDictionary[100]
            InputResultsTable["Details of ESI - Media Encryption"] = InputDictionary[102]
            InputResultsTable["Acquisition of ESI - Acquired By"] = InputDictionary[117]
            InputResultsTable["Acquisition of ESI - Date/Time Acquired"] = InputDictionary[128]
            InputResultsTable["Acquisition of ESI - Date/Time of Device"] = InputDictionary[135]
            InputResultsTable["Acquisition of ESI - Hardware Used - Tableau TD1"] = InputDictionary[152]
            InputResultsTable["Acquisition of ESI - Hardware Used - Tableau TD2"] = InputDictionary[154]
            InputResultsTable["Acquisition of ESI - Hardware Used - Tableau TD3"] = InputDictionary[156]
            InputResultsTable["Acquisition of ESI - Hardware Used - Cellebrite"] = InputDictionary[158]
            InputResultsTable["Acquisition of ESI - Hardware Used - Other"] = InputDictionary[161]
            InputResultsTable["Acquisition of ESI - Software Used - FTK"] = InputDictionary[164]
            InputResultsTable["Acquisition of ESI - Software Used - Encase"] = InputDictionary[166]
            InputResultsTable["Acquisition of ESI - Software Used - Solo"] = InputDictionary[168]
            InputResultsTable["Acquisition of ESI - Software Used - ExMerge"] = InputDictionary[170]
            InputResultsTable["Acquisition of ESI - Software Used - Cellebrite"] = InputDictionary[173]
            InputResultsTable["Acquisition of ESI - Software Used - Robocopy"] = InputDictionary[176]
            InputResultsTable["Acquisition of ESI - Software Used - Other"] = InputDictionary[178]
            InputResultsTable["Method of Acquisition - Physical"] = InputDictionary[188]
            InputResultsTable["Method of Acquisition - Logical"] = InputDictionary[190]
            InputResultsTable["Method of Acquisition - Targeted"] = InputDictionary[192]
            InputResultsTable["Method of Acquisition - Backup"] = InputDictionary[194]
            InputResultsTable["Method of Acquisition - Other"] = InputDictionary[197]
            InputResultsTable["Acquisition of ESI - Acquired Format - E01"] = InputDictionary[200]
            InputResultsTable["Acquisition of ESI - Acquired Format - AD1"] = InputDictionary[202]
            InputResultsTable["Acquisition of ESI - Acquired Format - RAW"] = InputDictionary[204]
            InputResultsTable["Acquisition of ESI - Acquired Format - DD"] = InputDictionary[206]
            InputResultsTable["Acquisition of ESI - Acquired Format - Other"] = InputDictionary[209]
            InputResultsTable["Acquisition of ESI - Acquired Size"] = InputDictionary[212]
            InputResultsTable["Acquisition of ESI - Image Verified"] = InputDictionary[220]
            InputResultsTable["Acquisition of ESI - Verified MD5 Hash"] = InputDictionary[224]
            InputResultsTable["Destination Media - Targeted Media"] = InputDictionary[243]
            InputResultsTable["Destination Media - Targeted Serial"] = InputDictionary[244]
            InputResultsTable["Destination Media - Backup Media"] = InputDictionary[245]
            InputResultsTable["Destination Media - Backup Serial"] = InputDictionary[246]
            InputResultsTable["Filename"] = file

            # Create empty Pandas Dataframe
            extractionframe = pd.DataFrame(columns=["Evidence Number",
                                                    "Case Information - Date",
                                                    "Case Information - Custodian",
                                                    "Case Information - Matter Name",
                                                    "Case Information - Matter Number",
                                                    "Case Information - Address Location",
                                                    "Type of ESI - Received By",
                                                    "Type of ESI - Desktop",
                                                    "Type of ESI - Laptop",
                                                    "Type of ESI - Server",
                                                    "Type of ESI - Tablet",
                                                    "Type of ESI - Mobile Device",
                                                    "Type of ESI - Hard Disk Drive",
                                                    "Type of ESI - Optical Disc",
                                                    "Type of ESI - Removable Media",
                                                    "Type of ESI - Removable Media Other",
                                                    "Type of ESI - Additional Details",
                                                    "Details of ESI - Device Make",
                                                    "Details of ESI - Media Make",
                                                    "Details of ESI - Device Model",
                                                    "Details of ESI - Media Model",
                                                    "Details of ESI - Device Serial",
                                                    "Details of ESI - Media Serial",
                                                    "Details of ESI - Device Asset Tag",
                                                    "Details of ESI - Media Encryption",
                                                    "Acquisition of ESI - Acquired By",
                                                    "Acquisition of ESI - Date/Time Acquired",
                                                    "Acquisition of ESI - Date/Time of Device",
                                                    "Acquisition of ESI - Hardware Used - Tableau TD1",
                                                    "Acquisition of ESI - Hardware Used - Tableau TD2",
                                                    "Acquisition of ESI - Hardware Used - Tableau TD3",
                                                    "Acquisition of ESI - Hardware Used - Cellebrite",
                                                    "Acquisition of ESI - Hardware Used - Other",
                                                    "Acquisition of ESI - Software Used - FTK",
                                                    "Acquisition of ESI - Software Used - Encase",
                                                    "Acquisition of ESI - Software Used - Solo",
                                                    "Acquisition of ESI - Software Used - ExMerge",
                                                    "Acquisition of ESI - Software Used - Cellebrite",
                                                    "Acquisition of ESI - Software Used - Robocopy",
                                                    "Acquisition of ESI - Software Used - Other",
                                                    "Method of Acquisition - Physical",
                                                    "Method of Acquisition - Logical",
                                                    "Method of Acquisition - Targeted",
                                                    "Method of Acquisition - Backup",
                                                    "Method of Acquisition - Other",
                                                    "Acquisition of ESI - Acquired Format - E01",
                                                    "Acquisition of ESI - Acquired Format - AD1",
                                                    "Acquisition of ESI - Acquired Format - RAW",
                                                    "Acquisition of ESI - Acquired Format - DD",
                                                    "Acquisition of ESI - Acquired Format - Other",
                                                    "Acquisition of ESI - Acquired Size",
                                                    "Acquisition of ESI - Image Verified",
                                                    "Acquisition of ESI - Verified MD5 Hash",
                                                    "Destination Media - Targeted Media",
                                                    "Destination Media - Targeted Serial",
                                                    "Destination Media - Backup Media",
                                                    "Destination Media - Backup Serial",
                                                    "Filename"
                                                    ])

            # Populate dataframe
            for values in InputResultsTable:
                extractionframe.loc[1, values] = InputResultsTable[values]

            print(extractionframe)

    except:
        messagebox.showerror("Error", "Please enter valid folder directory")

    try:
        os.chdir(working_directory)

        # Load Workbook
        workbook = load_workbook('TrackingSheet.xlsx')

        # Select Outputsheet
        output_sheet = workbook['01. Collection Details']

        # Append Results
        for row in dataframe_to_rows(extractionframe, index=False, header=False):
            output_sheet.append(row)

        workbook.save('TrackingSheet.xlsx')

    except:
        messagebox.showerror("Error", "Please enter valid working directory \n or check if TrackingSheet is present")

    tk.messagebox.showinfo("Progress Update", "Data Extracted, please exit")


def raise_frame(frame):
    frame.tkraise()


# initiate tinker, raise frame
master = tk.Tk()
f1 = Frame(master)

# write labels and entry boxes
Label(master, text="Directory with Tracking Sheet", font="Helvetica 8 bold italic").grid(row=1, column=1, padx=5,
                                                                                         pady=2)
Label(master, text="Directory with Folder Containing COCs", font="Helvetica 8 bold italic").grid(row=2, column=1,
                                                                                                 padx=5, pady=2)

e1 = Entry(master)
e2 = Entry(master)

e1.grid(row=1, column=2, padx=5, pady=2)
e2.grid(row=2, column=2, padx=5, pady=2)

# write submission and quit boxes
Button(master, text='Submit & Run Job', command=first_button).grid(row=4, column=1, columnspan=2, sticky=N, padx=5,
                                                                   pady=(5, 2))
Button(master, text='Exit', command=master.quit).grid(row=5, column=1, columnspan=2, sticky=N, padx=5, pady=(5, 2))

raise_frame(f1)

master.mainloop()